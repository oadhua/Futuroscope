import pandas as pd
import numpy as np
import warnings
from openpyxl import load_workbook
import re
import os

warnings.simplefilter(action='ignore', category=UserWarning)

# Danh sách các năm bạn cần xử lý
years = [2022, 2023, 2024, 2025, 2026] 
base_input_dir = r"D:\Stage SI\Machine Learning\Futuroscope\windows\donne_brut\cadence"
output_dir = r"D:\Stage SI\Machine Learning\Futuroscope\windows\donne_clean\cadence"

months = [
    "Janvier", "Fevrier", "Mars", "Avril", "Mai", "Juin", 
    "Juillet", "Aout", "Septembre", "Octobre", "Novembre", "Décembre"
]
target_rows = ["Jour", "Date", "JF", "H.  Ouv.", "H. Ferm", "Fréquentation", "Type Fréquentation"]

def remove_accents_and_special_chars(text):
    if not isinstance(text, str):
        return text
    text = text.lower().strip()
    replacements = {
        'à': 'a', 'á': 'a', 'â': 'a', 'ã': 'a', 'ä': 'a', 'å': 'a',
        'è': 'e', 'é': 'e', 'ê': 'e', 'ë': 'e',
        'ì': 'i', 'í': 'i', 'î': 'i', 'ï': 'i',
        'ò': 'o', 'ó': 'o', 'ô': 'o', 'õ': 'o', 'ö': 'o',
        'ù': 'u', 'ú': 'u', 'û': 'u', 'ü': 'u',
        'ç': 'c', 'ñ': 'n', 'ý': 'y', 'ÿ': 'y'
    }
    for accented, unaccented in replacements.items():
        text = text.replace(accented, unaccented)
    text = re.sub(r'[\s\.\-\(\)\[\]]+', '_', text)
    return text.strip('_')

all_years_dfs = []

for year in years:
    file_path = os.path.join(base_input_dir, f"horaire{year}.xlsx")
    if not os.path.exists(file_path):
        print(f"Bỏ qua năm {year} do không tìm thấy file: {file_path}")
        continue
        
    print(f"\n=== ĐANG XỬ LÝ NĂM {year} ===")
    wb = load_workbook(file_path, data_only=True)
    all_months_data = []

    for month in months:
        try:
            if month not in wb.sheetnames:
                continue
            ws = wb[month]
            
            merged_cells_dict = {}
            for merged_range in ws.merged_cells.ranges:
                top_left_cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
                val = top_left_cell.value
                for r in range(merged_range.min_row, merged_range.max_row + 1):
                    for c in range(merged_range.min_col, merged_range.max_col + 1):
                        merged_cells_dict[(r, c)] = val

            grid_data = []
            for r in range(1, ws.max_row + 1):
                row_vals = []
                for c in range(1, ws.max_column + 1):
                    row_vals.append(merged_cells_dict.get((r, c), ws.cell(row=r, column=c).value))
                grid_data.append(row_vals)

            df_raw = pd.DataFrame(grid_data)
            df_filtered = df_raw[df_raw[0].isin(target_rows)].copy()
            if df_filtered.empty:
                continue

            df_transposed = df_filtered.set_index(0).T.dropna(subset=["Date"])
            df_transposed = df_transposed[df_transposed["Date"].astype(str).str.contains(r'\d{4}')]

            if not df_transposed.empty:
                df_transposed["Date"] = pd.to_datetime(df_transposed["Date"])
                df_transposed["Nom_Jour"] = df_transposed["Date"].dt.day_name()
                df_transposed["Jour"] = df_transposed["Date"].dt.day
                df_transposed["Mois"] = df_transposed["Date"].dt.month
                df_transposed["Annee"] = df_transposed["Date"].dt.year
                df_transposed["is_weekend"] = df_transposed["Date"].dt.dayofweek.isin([5, 6]).astype(int)
                
                # Tìm cột H. Ouv. bất kể khoảng trắng thừa
                ouv_col = [c for c in df_transposed.columns if 'ouv' in str(c).lower()]
                ferm_col = [c for c in df_transposed.columns if 'ferm' in str(c).lower()]

                if ouv_col and ferm_col:
                    has_ouv = df_transposed[ouv_col[0]].notna() & (df_transposed[ouv_col[0]].astype(str).str.strip() != "")
                    has_ferm = df_transposed[ferm_col[0]].notna() & (df_transposed[ferm_col[0]].astype(str).str.strip() != "")
                    df_transposed["is_open"] = (has_ouv & has_ferm).astype(int)
                else:
                    df_transposed["is_open"] = 0
            
            df_transposed.insert(0, "Mois_Nom", month)
            all_months_data.append(df_transposed)
        except Exception as e:
            print(f"Lỗi tại tab {month} năm {year}: {e}")

    if all_months_data:
        df_year = pd.concat(all_months_data, ignore_index=True, join='outer')
        all_years_dfs.append(df_year)

# GỘP TẤT CẢ CÁC NĂM THÀNH 1 FILE CSV DUY NHẤT
if all_years_dfs:
    final_df = pd.concat(all_years_dfs, ignore_index=True, join='outer')
    
    if "JF" in final_df.columns:
        final_df["JF"] = np.where(
            final_df["JF"].isna() | 
            (final_df["JF"].astype(str).str.strip() == '') | 
            (final_df["JF"].astype(str).str.strip() == '0'),
            0, 1
        ).astype(int)

    # CHỈNH SỬA CHÍNH: Chuyển Date sang chuẩn ISO YYYY-MM-DD để đồng bộ với các file khác
    if "Date" in final_df.columns:
        final_df["Date"] = pd.to_datetime(final_df["Date"]).dt.strftime('%Y-%m-%d')

    # Tìm lại chính xác tên cột H. Ouv. và H. Ferm trong final_df
    col_map = {}
    for col in final_df.columns:
        if 'ouv' in str(col).lower():
            col_map[col] = "H_Ouv"
        elif 'ferm' in str(col).lower():
            col_map[col] = "H_Ferm"
    final_df = final_df.rename(columns=col_map)

    ordered_cols = [
        "Mois_Nom", "Date", "JF", "Nom_Jour", "Jour", "Mois", "Annee", "is_weekend", "is_open",
        "H_Ouv", "H_Ferm", "Fréquentation", "Type Fréquentation"
    ]
    final_cols = [col for col in ordered_cols if col in final_df.columns]
    final_df = final_df[final_cols]
    final_df.columns = [remove_accents_and_special_chars(col) for col in final_df.columns]
    
    # Xuất ra file CSV duy nhất
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "horaire_all_years.csv")
    final_df.to_csv(output_path, index=False, encoding='utf-8-sig')
    print(f"\n===> ĐÃ TẠO FILE GỘP THÀNH CÔNG: {output_path}")